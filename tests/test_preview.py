import subprocess
from pathlib import Path
from zipfile import ZipFile

import pytest

from core.compression import ArchiveFormat, SevenZipCompressionService, find_7zip_executable
from core.preview import PreviewService


def test_preview_service_handles_image_text_pdf_office_archive_and_video(tmp_path: Path):
    service = PreviewService(tmp_path / "cache")

    from PIL import Image
    image_path = tmp_path / "photo.png"
    Image.new("RGB", (320, 180), "red").save(image_path)
    image = service.preview(image_path)
    assert image.kind == "image"
    assert image.image_path and Path(image.image_path).exists()
    assert image.metadata["dimensions"] == "320 x 180"

    text_path = tmp_path / "notes.txt"
    text_path.write_text("hello preview", encoding="utf-8")
    text = service.preview(text_path)
    assert text.kind == "text"
    assert "hello preview" in text.text

    import fitz
    pdf_path = tmp_path / "assignment.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "PDF preview text")
    doc.save(pdf_path)
    doc.close()
    pdf = service.preview(pdf_path)
    assert pdf.kind == "pdf"
    assert pdf.image_path and Path(pdf.image_path).exists()
    assert "PDF preview text" in pdf.text

    import docx
    docx_path = tmp_path / "assignment.docx"
    word = docx.Document()
    word.add_paragraph("DOCX preview text")
    word.save(docx_path)
    word_preview = service.preview(docx_path)
    assert word_preview.kind == "docx"
    assert "DOCX preview text" in word_preview.text

    import openpyxl
    xlsx_path = tmp_path / "sheet.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Name"
    ws["B1"] = "Value"
    wb.save(xlsx_path)
    xlsx = service.preview(xlsx_path)
    assert xlsx.kind == "xlsx"
    assert "Data" in xlsx.text

    from pptx import Presentation
    pptx_path = tmp_path / "slides.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = "Slide title"
    deck.save(pptx_path)
    pptx = service.preview(pptx_path)
    assert pptx.kind == "pptx"
    assert "Slide title" in pptx.text

    zip_path = tmp_path / "archive.zip"
    with ZipFile(zip_path, "w") as archive:
        archive.writestr("top/file.txt", "content")
    archive_preview = service.preview(zip_path)
    assert archive_preview.kind == "archive"
    assert "top/file.txt" in archive_preview.text
    assert archive_preview.metadata["entries"] == [
        {"path": "top/file.txt", "size_bytes": 7, "is_dir": False}
    ]

    import imageio_ffmpeg
    video_path = tmp_path / "clip.mp4"
    subprocess.run(
        [
            imageio_ffmpeg.get_ffmpeg_exe(),
            "-y",
            "-f",
            "lavfi",
            "-i",
            "color=c=blue:s=160x90:d=1",
            str(video_path),
        ],
        check=True,
        capture_output=True,
    )
    video = service.preview(video_path)
    assert video.kind == "video"
    assert video.image_path and Path(video.image_path).exists()


def test_preview_service_lists_7z_as_a_container(tmp_path: Path):
    executable = find_7zip_executable()
    if executable is None:
        pytest.skip("7-Zip is not installed or bundled")
    source = tmp_path / "source"
    (source / "src").mkdir(parents=True)
    (source / "src" / "app.py").write_text("print('Jena')", encoding="utf-8")
    archive = tmp_path / "source.7z"
    engine = SevenZipCompressionService(executable)
    engine.create(source, ["src/app.py"], archive, ArchiveFormat.SEVEN_ZIP)

    preview = PreviewService(tmp_path / "cache", compression=engine).preview(archive)

    assert preview.kind == "archive"
    assert preview.metadata["file_count"] == 1
    assert preview.metadata["entries"][0]["path"] == "src/app.py"
